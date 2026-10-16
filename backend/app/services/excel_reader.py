"""Read batch codes from an .xlsx upload.

RF-09 (first sheet), RF-10 (column detection), RF-11 (numeric/scientific values and
leading zeros), RF-18 (empty rows), RF-19 (limits), RNF-03 (size first),
RNF-05 (read_only + data_only), RNF-11 (content signature, no macros).
"""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.core.errors import AppError, ErrorCode

BYTES_PER_MB = 1024 * 1024
FULL_LENGTH = 13

_ZIP_SIGNATURE = b"PK\x03\x04"
_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0"  # legacy .xls (and other old Office files)
# Zip bomb guard: a 2 MB upload must not expand into gigabytes in 512 MB of RAM.
_MAX_UNCOMPRESSED_BYTES = 100 * BYTES_PER_MB
# A sheet can carry a million formatted but empty rows; don't walk them all.
_MAX_SCANNED_ROWS = 100_000
_HEADER_SEARCH_ROWS = 5

_CODE_HEADERS = {"codigo", "ean", "ean13", "gtin"}
_NAME_HEADERS = {"nombrearchivo", "nombre", "archivo"}

_SCIENTIFIC_TEXT = re.compile(r"[+]?(\d+)(?:[.,](\d+))?[eE]\+?(\d+)")
_SCIENTIFIC_LOST_ZEROS = 3

_TEXT_FORMAT_HINT = "Formatea la columna «codigo» como Texto y vuelve a escribir el código."


@dataclass(frozen=True, slots=True)
class SheetRow:
    row: int
    """Excel row number, as the user sees it."""
    input: str
    """Original cell value as text, for the report."""
    code: str | None
    """Text to validate with the symbology; None when the cell can't be a code."""
    error: str | None
    """Why the cell can't be read as a code (RF-11)."""
    name: str | None
    """Raw ``nombre_archivo`` value, not sanitized yet."""


def read_sheet(
    content: bytes, *, max_rows: int, max_mb: int, restore_leading_zeros: bool
) -> list[SheetRow]:
    """Rows with a code or a name from the first sheet; empty rows are skipped."""
    check_file(content, max_mb=max_mb)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many exception types for corrupt files
        raise AppError(
            ErrorCode.INVALID_FILE_TYPE,
            "No se pudo leer el archivo. Verifica que sea un Excel .xlsx válido.",
        ) from exc

    try:
        rows = enumerate(workbook.worksheets[0].iter_rows(), start=1)
        code_column, name_column = _find_columns(rows)
        result: list[SheetRow] = []
        for row_number, cells in rows:
            if row_number > _MAX_SCANNED_ROWS:
                raise _too_many_rows(max_rows)
            row = _read_row(row_number, cells, code_column, name_column, restore_leading_zeros)
            if row is None:
                continue  # RF-18: empty rows are not errors
            result.append(row)
            if len(result) > max_rows:
                raise _too_many_rows(max_rows)
        return result
    finally:
        workbook.close()


def check_file(content: bytes, *, max_mb: int) -> None:
    """Size and content checks before any parsing (RNF-03, RNF-11)."""
    if len(content) > max_mb * BYTES_PER_MB:
        raise AppError(
            ErrorCode.FILE_TOO_LARGE,
            f"El archivo supera el límite de {max_mb} MB.",
            {"max_mb": max_mb},
        )
    if content.startswith(_OLE_SIGNATURE):
        raise _invalid_file("El formato .xls antiguo no está soportado. Guárdalo como .xlsx.")
    if not content.startswith(_ZIP_SIGNATURE):
        raise _invalid_file("El archivo no es un Excel .xlsx válido.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            content_types = (
                archive.read("[Content_Types].xml") if "[Content_Types].xml" in names else b""
            )
    except zipfile.BadZipFile as exc:
        raise _invalid_file("El archivo no es un Excel .xlsx válido.") from exc

    if "xl/workbook.xml" not in names or not content_types:
        raise _invalid_file("El archivo no es un Excel .xlsx válido.")
    if any(name.endswith("vbaProject.bin") for name in names) or b"macroEnabled" in content_types:
        raise _invalid_file("Los archivos con macros (.xlsm) no están permitidos.")
    if sum(entry.file_size for entry in entries) > _MAX_UNCOMPRESSED_BYTES:
        raise _invalid_file("El archivo es demasiado grande al descomprimirse.")


def _find_columns(rows: Any) -> tuple[int, int | None]:
    """Locate the header (first rows) and return the code and name column indexes."""
    first_row: list[str] = []
    for row_number, cells in rows:
        headers = [_header_key(cell.value) for cell in cells]
        code_column = next((i for i, h in enumerate(headers) if h in _CODE_HEADERS), None)
        if code_column is not None:
            name_column = next((i for i, h in enumerate(headers) if h in _NAME_HEADERS), None)
            return code_column, name_column
        if not first_row:  # what the user probably meant as headers
            first_row = [str(c.value).strip() for c in cells if c.value not in (None, "")]
        if row_number >= _HEADER_SEARCH_ROWS:
            break
    raise AppError(
        ErrorCode.CODE_COLUMN_NOT_FOUND,
        "No se encontró la columna de códigos. Usa un encabezado «codigo», «ean», «ean13» "
        "o «gtin» en la primera fila.",
        {"headers_found": first_row[:20]},
    )


def _header_key(value: object) -> str:
    """'Código', ' EAN-13 ', 'Nombre_Archivo' → 'codigo', 'ean13', 'nombrearchivo'."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _read_row(
    row_number: int,
    cells: tuple[Any, ...],
    code_column: int,
    name_column: int | None,
    restore_leading_zeros: bool,
) -> SheetRow | None:
    code_cell = cells[code_column] if code_column < len(cells) else None
    value = code_cell.value if code_cell is not None else None
    number_format = str(getattr(code_cell, "number_format", None) or "General")
    name = None
    if name_column is not None and name_column < len(cells):
        name = _text(cells[name_column].value) or None

    if _is_blank(value):
        if name is None:
            return None
        return SheetRow(row_number, "", None, "Falta el código.", name)

    input_text, code, error = _normalize_code(value, number_format, restore_leading_zeros)
    return SheetRow(row_number, input_text, code, error, name)


def _normalize_code(
    value: object, number_format: str, restore_leading_zeros: bool
) -> tuple[str, str | None, str | None]:
    """Return (original text, code to validate, error) for one code cell (RF-11)."""
    if isinstance(value, bool):
        return str(value), None, "El valor no es un código."
    if isinstance(value, datetime | date | time):
        return (
            value.isoformat(),
            None,
            f"Excel interpretó el valor como una fecha. {_TEXT_FORMAT_HINT}",
        )
    if isinstance(value, str):
        text = value.strip()
        if match := _SCIENTIFIC_TEXT.fullmatch(text):
            return _from_scientific_text(text, match)
        return value, text, None
    if isinstance(value, int | float):
        return _from_number(value, number_format, restore_leading_zeros)
    return str(value), None, "El valor no es un código."


def _from_number(
    value: float, number_format: str, restore_leading_zeros: bool
) -> tuple[str, str | None, str | None]:
    if value < 0 or (isinstance(value, float) and not value.is_integer()):
        return _text(value), None, "El código debe ser un número entero de 12 o 13 dígitos."
    digits = str(int(value))

    # Excel keeps the full number, but if the cell is shown in scientific notation and
    # ends in zeros, it was most likely typed (or pasted) already truncated.
    if "E" in number_format.upper() and digits.endswith("0" * _SCIENTIFIC_LOST_ZEROS):
        return (
            digits,
            None,
            f"El código está en notación científica y pudo perder dígitos. {_TEXT_FORMAT_HINT}",
        )

    # Numbers lose leading zeros: 0012345678905 arrives as 12345678905.
    if len(digits) in (11, 12) and restore_leading_zeros:
        return digits, digits.zfill(FULL_LENGTH), None
    if len(digits) == 11:
        return (
            digits,
            None,
            "Tiene 11 dígitos: si Excel quitó ceros iniciales, activa «Restaurar ceros "
            f"iniciales» o {_TEXT_FORMAT_HINT[0].lower()}{_TEXT_FORMAT_HINT[1:]}",
        )
    return digits, digits, None


def _from_scientific_text(text: str, match: re.Match[str]) -> tuple[str, str | None, str | None]:
    """'7.86123E+12' typed or imported as text."""
    integer, fraction, _ = match.groups()
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation:  # pragma: no cover - the regex only lets valid numbers through
        return text, None, "El valor no es un código."
    if number != number.to_integral_value():
        return text, None, "El código debe ser un número entero de 12 o 13 dígitos."
    digits = str(int(number))
    significant = len((integer + (fraction or "")).lstrip("0"))
    if significant < len(digits):
        return (
            text,
            None,
            f"Excel convirtió el código a notación científica ({text}) y se perdieron "
            f"dígitos. {_TEXT_FORMAT_HINT}",
        )
    return text, digits, None


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _invalid_file(message: str) -> AppError:
    return AppError(ErrorCode.INVALID_FILE_TYPE, message)


def _too_many_rows(max_rows: int) -> AppError:
    return AppError(
        ErrorCode.TOO_MANY_ROWS,
        f"El archivo supera el límite de {max_rows} filas con códigos.",
        {"max_rows": max_rows},
    )
