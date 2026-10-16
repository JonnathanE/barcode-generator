"""Downloadable Excel template for batch generation (RF-08)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TEMPLATE_FILENAME = "plantilla_codigos.xlsx"
TEXT_FORMAT = "@"
# Pre-formatted as Text so Excel keeps leading zeros and never shows 7.86E+12.
FORMATTED_ROWS = 1_000

EXAMPLES = [
    ("786123456789", "producto-ejemplo-1"),
    ("7861234567898", "producto-ejemplo-2"),
    ("0012345678905", ""),
]

INSTRUCTIONS = [
    "Cómo usar esta plantilla",
    "",
    "1. Escribe un código por fila en la columna «codigo» de la hoja «Códigos».",
    "   - 12 dígitos: se calcula el dígito de control.",
    "   - 13 dígitos: se verifica el dígito de control.",
    "   - Se ignoran espacios y guiones.",
    "2. «nombre_archivo» es opcional. Si está vacío, el archivo se llama como el código.",
    "3. Borra las filas de ejemplo antes de subir el archivo.",
    "4. La columna «codigo» ya está en formato Texto: así Excel no quita ceros iniciales",
    "   ni convierte los códigos a notación científica (7.86123E+12).",
    "",
    "Los códigos EAN-13 deben estar asignados por GS1 (en Ecuador, GS1 Ecuador).",
]


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Códigos"

    sheet.append(["codigo", "nombre_archivo"])
    header_fill = PatternFill("solid", fgColor="E5E7EB")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row, (code, name) in enumerate(EXAMPLES, start=2):
        sheet.cell(row=row, column=1, value=code)
        sheet.cell(row=row, column=2, value=name or None)

    for row in range(2, FORMATTED_ROWS + 2):
        sheet.cell(row=row, column=1).number_format = TEXT_FORMAT
        sheet.cell(row=row, column=2).number_format = TEXT_FORMAT

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 32
    sheet.freeze_panes = "A2"

    help_sheet = workbook.create_sheet("Instrucciones")
    for line in INSTRUCTIONS:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True, size=13)
    help_sheet.column_dimensions["A"].width = 90
    for row in help_sheet.iter_rows():
        row[0].alignment = Alignment(wrap_text=False)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
