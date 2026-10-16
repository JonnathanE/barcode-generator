import io
import json
import logging
import os
import shutil
import zipfile
from pathlib import Path
from typing import Any

import pytest
import zxingcpp
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.core.config import get_settings
from app.main import app
from tests.disk_writes import no_disk_writes

client = TestClient(app, raise_server_exceptions=False)
FIXTURES = Path(__file__).parent.parent / "fixtures"
needs_gs = pytest.mark.skipif(shutil.which("gs") is None, reason="Ghostscript no instalado")
VALIDATE = "/api/v1/batch/ean13/validate"
GENERATE = "/api/v1/batch/ean13/generate"


@pytest.fixture(autouse=True)
def generous_batch_rate_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "rate_limit_batch", "1000/minute")


def upload(name: str, content: bytes | None = None) -> dict[str, Any]:
    data = content if content is not None else (FIXTURES / name).read_bytes()
    return {"file": (name, data, "application/octet-stream")}


def validate(name: str, **form: Any) -> Any:
    return client.post(VALIDATE, files=upload(name), data=form)


def statuses(body: dict[str, Any]) -> list[tuple[int, str]]:
    return [(row["row"], row["status"]) for row in body["rows"]]


def workbook(rows: list[list[Any]]) -> bytes:
    book = Workbook()
    for row in rows:
        book.worksheets[0].append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


# ---------- Template ----------


def test_template_download() -> None:
    response = client.get("/api/v1/batch/template")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        response.headers["content-disposition"] == 'attachment; filename="plantilla_codigos.xlsx"'
    )
    assert load_workbook(io.BytesIO(response.content)).sheetnames == ["Códigos", "Instrucciones"]


# ---------- Step 1: validate, one test per fixture ----------


def test_validos() -> None:
    body = validate("validos.xlsx").json()

    assert body["summary"] == {"total": 4, "valid": 4, "errors": 0, "duplicates": 0}
    assert [r["normalized"] for r in body["rows"]] == [
        "7861234567898",
        "9780306406157",
        "4006381333931",
        "5901234123457",
    ]
    assert body["rows"][0] == {
        "row": 2,
        "input": "786123456789",
        "normalized": "7861234567898",
        "status": "ok",
        "message": "",
        "filename": "leche-entera",
    }


def test_invalidos() -> None:
    body = validate("invalidos.xlsx").json()

    assert body["summary"] == {"total": 5, "valid": 1, "errors": 4, "duplicates": 0}
    messages = [row["message"] for row in body["rows"][:4]]
    assert messages == [
        "El código solo puede contener dígitos (se permiten espacios y guiones).",
        "El código debe tener 12 o 13 dígitos; tiene 5.",
        "El dígito de control debe ser 8.",
        "Falta el código.",
    ]


def test_numericos_without_and_with_restore_flag() -> None:
    plain = validate("numericos.xlsx").json()
    restored = validate("numericos.xlsx", restore_leading_zeros="true").json()

    assert statuses(plain) == [(2, "ok"), (3, "ok"), (4, "ok"), (5, "error"), (6, "error")]
    # With the flag, the 12-digit body becomes 0786123456789, whose check digit is wrong:
    # the check digit catches a wrongly applied flag.
    assert statuses(restored) == [(2, "error"), (3, "ok"), (4, "ok"), (5, "ok"), (6, "error")]
    assert restored["rows"][3]["normalized"] == "0012345678905"


def test_cientifica() -> None:
    body = validate("cientifica.xlsx").json()

    assert statuses(body) == [(2, "error"), (3, "ok"), (4, "error"), (5, "duplicado")]
    assert "se perdieron dígitos" in body["rows"][0]["message"]


def test_duplicados() -> None:
    body = validate("duplicados.xlsx").json()

    assert body["summary"] == {"total": 4, "valid": 2, "errors": 0, "duplicates": 2}
    assert statuses(body) == [(2, "ok"), (3, "duplicado"), (4, "duplicado"), (5, "ok")]


def test_filas_vacias() -> None:
    body = validate("filas_vacias.xlsx").json()

    assert statuses(body) == [(3, "ok"), (7, "ok")]


def test_nombres_peligrosos() -> None:
    body = validate("nombres_peligrosos.xlsx").json()

    assert [r["filename"] for r in body["rows"]] == [
        "___x",
        "___windows_system32",
        "_CON",
        "a_b_c_d",
        "Leche",
        "leche_2",
        "7860000000072",
        "x" * 100,
    ]


def test_encabezado_con_titulo() -> None:
    assert statuses(validate("encabezado_con_titulo.xlsx").json()) == [(4, "ok")]


@pytest.mark.parametrize(
    ("name", "status", "code"),
    [
        ("sin_columna_codigo.xlsx", 422, "CODE_COLUMN_NOT_FOUND"),
        ("no_es_excel.xlsx", 415, "INVALID_FILE_TYPE"),
        ("zip_no_excel.xlsx", 415, "INVALID_FILE_TYPE"),
        ("con_macros.xlsx", 415, "INVALID_FILE_TYPE"),
    ],
)
def test_rejected_files(name: str, status: int, code: str) -> None:
    response = validate(name)

    assert response.status_code == status
    assert response.json()["error"]["code"] == code


def test_lote_500_fits_the_default_limit() -> None:
    assert validate("lote_500.xlsx").json()["summary"]["valid"] == 500


def test_too_many_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "max_batch_rows", 3)

    response = validate("validos.xlsx")

    assert response.status_code == 422
    assert (
        response.json()["error"]["message"] == "El archivo supera el límite de 3 filas con códigos."
    )


# ---------- Upload size and storage ----------


def test_oversized_upload_is_rejected_before_parsing_with_cors() -> None:
    too_big = b"PK\x03\x04" + os.urandom(3 * 1024 * 1024)

    response = client.post(
        VALIDATE, files=upload("grande.xlsx", too_big), headers={"Origin": "http://localhost:5173"}
    )

    assert response.status_code == 413
    assert response.json()["error"] == {
        "code": "FILE_TOO_LARGE",
        "message": "El archivo supera el límite de 2 MB.",
        "details": {"max_mb": 2},
    }
    assert response.headers["access-control-allow-origin"] == "http://localhost:5173"


def test_upload_without_content_length_is_rejected() -> None:
    def chunks() -> Any:
        yield b"--x\r\n"

    response = client.post(
        VALIDATE, content=chunks(), headers={"Content-Type": "multipart/form-data; boundary=x"}
    )

    assert response.status_code == 411


def test_uploads_near_the_limit_stay_in_memory() -> None:
    content = (FIXTURES / "validos.xlsx").read_bytes()
    padded = io.BytesIO(content)
    with zipfile.ZipFile(padded, "a", zipfile.ZIP_STORED) as archive:
        archive.writestr("xl/media/relleno.bin", os.urandom(1_800_000))  # > Starlette's 1 MB
    assert 1_048_576 < len(padded.getvalue()) < 2 * 1_048_576

    with no_disk_writes() as writes:
        response = client.post(VALIDATE, files=upload("casi2mb.xlsx", padded.getvalue()))

    assert response.status_code == 200
    assert writes == []


# ---------- Step 2: generate ----------


@needs_gs
def test_generate_returns_zip_with_both_formats() -> None:
    response = client.post(
        GENERATE,
        files=upload("validos.xlsx"),
        data={"formats": ["png", "eps"], "options": json.dumps({"scale": 1, "dpi": 150})},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="codigos_') and disposition.endswith(
        '.zip"'
    )
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    names = archive.namelist()
    assert len([n for n in names if n.startswith("png/")]) == 4
    assert len([n for n in names if n.startswith("eps/")]) == 4
    assert "reporte.csv" in names
    image = Image.open(io.BytesIO(archive.read("png/leche-entera.png")))
    assert [r.text for r in zxingcpp.read_barcodes(image)] == ["7861234567898"]
    assert image.info["dpi"] == pytest.approx((150, 150), abs=0.05)


@needs_gs
def test_generate_only_png_with_restore_flag() -> None:
    response = client.post(
        GENERATE,
        files=upload("numericos.xlsx"),
        data={"formats": "png", "restore_leading_zeros": "true"},
    )

    names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
    assert sorted(names) == [
        "png/0012345678905.png",
        "png/4006381333931.png",
        "png/5901234123457.png",
        "reporte.csv",
    ]


@pytest.mark.parametrize(
    ("data", "field"),
    [
        ({"formats": "svg"}, "body.formats.0"),
        ({}, "body.formats"),
    ],
)
def test_generate_rejects_bad_formats(data: dict[str, Any], field: str) -> None:
    response = client.post(GENERATE, files=upload("validos.xlsx"), data=data)

    assert response.status_code == 422
    error = response.json()["error"]
    assert error["code"] == "INVALID_REQUEST"
    assert field in [f["field"] for f in error["details"]["fields"]]


@pytest.mark.parametrize(
    ("options", "field"),
    [('{"scale": 9}', "options.scale"), ("no-json", "options")],
)
def test_generate_rejects_bad_options(options: str, field: str) -> None:
    response = client.post(
        GENERATE, files=upload("validos.xlsx"), data={"formats": "png", "options": options}
    )

    assert response.status_code == 422
    error = response.json()["error"]
    assert error["message"] == "Las opciones de generación no son válidas."
    assert field in [f["field"] for f in error["details"]["fields"]]


def test_generate_without_valid_rows() -> None:
    content = workbook([["codigo"], ["123"], ["abc"]])

    response = client.post(GENERATE, files=upload("malos.xlsx", content), data={"formats": "png"})

    assert response.status_code == 422
    error = response.json()["error"]
    assert error["message"] == "El archivo no tiene códigos válidos para generar."
    assert error["details"]["summary"] == {"total": 2, "valid": 0, "errors": 2, "duplicates": 0}


# ---------- Rate limit and logs ----------


def test_batch_rate_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "rate_limit_batch", "2/minute")

    codes = [validate("validos.xlsx").status_code for _ in range(3)]

    assert codes == [200, 200, 429]


def test_batch_logs_counts_but_never_codes(caplog: pytest.LogCaptureFixture) -> None:
    with caplog.at_level(logging.INFO):
        validate("validos.xlsx")

    (record,) = [r for r in caplog.records if r.getMessage() == "batch_validated"]
    assert (record.total, record.valid) == (4, 4)  # type: ignore[attr-defined]
    logged = " ".join(str(vars(r)) for r in caplog.records)
    assert "7861234567898" not in logged and "leche-entera" not in logged
