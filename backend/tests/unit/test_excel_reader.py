import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errors import AppError, ErrorCode
from app.services import excel_reader
from app.services.excel_reader import SheetRow, check_file, read_sheet
from app.services.template import TEMPLATE_FILENAME, build_template

FIXTURES = Path(__file__).parent.parent / "fixtures"


def read(name: str, *, max_rows: int = 500, restore: bool = False) -> list[SheetRow]:
    content = (FIXTURES / name).read_bytes()
    return read_sheet(content, max_rows=max_rows, max_mb=2, restore_leading_zeros=restore)


def codes(rows: list[SheetRow]) -> list[str | None]:
    return [row.code for row in rows]


def expect_error(
    code: ErrorCode, name: str | None = None, content: bytes | None = None
) -> AppError:
    data = content if content is not None else (FIXTURES / str(name)).read_bytes()
    with pytest.raises(AppError) as exc_info:
        read_sheet(data, max_rows=500, max_mb=2, restore_leading_zeros=False)
    assert exc_info.value.code is code
    return exc_info.value


# ---------- Reading and column detection ----------


def test_valid_text_codes_keep_row_numbers_and_names() -> None:
    rows = read("validos.xlsx")

    assert [(r.row, r.code, r.name, r.error) for r in rows] == [
        (2, "786123456789", "leche-entera", None),
        (3, "9780306406157", None, None),
        (4, "400-6381-33393-1", "galletas", None),
        (5, "590 1234 12345 7", None, None),
    ]


def test_header_can_be_below_a_title_with_accents_and_case() -> None:
    rows = read("encabezado_con_titulo.xlsx")

    assert [(r.row, r.code, r.name) for r in rows] == [(4, "7861234567898", "Leche")]


def test_missing_code_column_lists_headers_found() -> None:
    error = expect_error(ErrorCode.CODE_COLUMN_NOT_FOUND, "sin_columna_codigo.xlsx")

    assert error.details == {"headers_found": ["producto", "precio"]}


def test_empty_rows_are_ignored() -> None:
    rows = read("filas_vacias.xlsx")

    assert [(r.row, r.code) for r in rows] == [(3, "7861234567898"), (7, "4006381333931")]


def test_row_with_name_but_no_code_is_an_error() -> None:
    rows = read("invalidos.xlsx")

    missing = next(r for r in rows if r.name == "sin-codigo")
    assert (missing.row, missing.code, missing.error) == (5, None, "Falta el código.")


def test_invalid_text_codes_are_left_to_the_symbology() -> None:
    rows = read("invalidos.xlsx")

    assert codes(rows) == ["78612345678A", "12345", "7861234567897", None, "4006381333931"]


# ---------- RF-11: numbers, scientific notation, leading zeros ----------


def test_numbers_are_read_as_digits() -> None:
    rows = read("numericos.xlsx")

    assert [(r.input, r.code) for r in rows[:3]] == [
        ("786123456789", "786123456789"),
        ("5901234123457", "5901234123457"),
        ("4006381333931", "4006381333931"),
    ]
    decimal = rows[4]
    assert (decimal.code, decimal.error) == (
        None,
        "El código debe ser un número entero de 12 o 13 dígitos.",
    )


def test_eleven_digit_number_is_an_error_without_restore_flag() -> None:
    row = read("numericos.xlsx")[3]

    assert row.code is None
    assert row.error is not None
    assert "11 dígitos" in row.error and "Restaurar ceros iniciales" in row.error


def test_restore_flag_pads_numbers_to_13_digits() -> None:
    rows = read("numericos.xlsx", restore=True)

    assert codes(rows)[:4] == ["0786123456789", "5901234123457", "4006381333931", "0012345678905"]


def test_restore_flag_does_not_touch_text_cells() -> None:
    assert codes(read("validos.xlsx", restore=True))[0] == "786123456789"


def test_scientific_notation() -> None:
    rows = read("cientifica.xlsx")

    lost_text, full_text, lost_number, full_number = rows
    assert lost_text.code is None
    assert lost_text.error is not None and "se perdieron dígitos" in lost_text.error
    assert lost_text.input == "7.86123E+12"
    assert full_text.code == "7861234567898"
    assert lost_number.code is None
    assert lost_number.error is not None and "notación científica" in lost_number.error
    assert full_number.code == "7861234567898"


@pytest.mark.parametrize(
    ("value", "expected_code", "expected_error"),
    [
        (True, None, "El valor no es un código."),
        (-5, None, "El código debe ser un número entero de 12 o 13 dígitos."),
        ("7,861234567898E+12", "7861234567898", None),
        ("1.5E+0", None, "El código debe ser un número entero de 12 o 13 dígitos."),
    ],
)
def test_unusual_cell_values(
    value: object, expected_code: str | None, expected_error: str | None
) -> None:
    _, code, error = excel_reader._normalize_code(value, "General", restore_leading_zeros=False)

    assert (code, error) == (expected_code, expected_error)


def test_dates_suggest_text_format() -> None:
    book = Workbook()
    sheet = book.worksheets[0]
    sheet.append(["codigo"])
    sheet.append([__import__("datetime").date(2026, 10, 16)])
    output = io.BytesIO()
    book.save(output)

    (row,) = read_sheet(output.getvalue(), max_rows=5, max_mb=2, restore_leading_zeros=False)

    assert row.error is not None and "fecha" in row.error


# ---------- RF-19 / RNF-03 / RNF-11: limits and file checks ----------


def test_too_many_rows_rejects_the_whole_file() -> None:
    with pytest.raises(AppError) as exc_info:
        read("lote_500.xlsx", max_rows=499)

    assert exc_info.value.code is ErrorCode.TOO_MANY_ROWS
    assert exc_info.value.message == "El archivo supera el límite de 499 filas con códigos."
    assert exc_info.value.details == {"max_rows": 499}


def test_max_rows_is_inclusive() -> None:
    assert len(read("lote_500.xlsx", max_rows=500)) == 500


def test_file_size_is_checked_before_parsing() -> None:
    oversized = b"PK\x03\x04" + b"0" * (2 * 1024 * 1024)

    error = expect_error(ErrorCode.FILE_TOO_LARGE, content=oversized)

    assert error.status_code == 413
    assert error.details == {"max_mb": 2}


@pytest.mark.parametrize(
    ("name", "message"),
    [
        ("no_es_excel.xlsx", "El archivo no es un Excel .xlsx válido."),
        ("zip_no_excel.xlsx", "El archivo no es un Excel .xlsx válido."),
        ("con_macros.xlsx", "Los archivos con macros (.xlsm) no están permitidos."),
    ],
)
def test_content_is_checked_not_the_extension(name: str, message: str) -> None:
    error = expect_error(ErrorCode.INVALID_FILE_TYPE, name)

    assert error.message == message


def test_legacy_xls_has_a_specific_message() -> None:
    error = expect_error(ErrorCode.INVALID_FILE_TYPE, content=b"\xd0\xcf\x11\xe0" + b"\0" * 100)

    assert ".xls antiguo" in error.message


def test_corrupt_zip_with_valid_signature() -> None:
    expect_error(ErrorCode.INVALID_FILE_TYPE, content=b"PK\x03\x04" + b"garbage" * 10)


def test_zip_bomb_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(excel_reader, "_MAX_UNCOMPRESSED_BYTES", 10_000)
    content = (FIXTURES / "validos.xlsx").read_bytes()
    output = io.BytesIO(content)
    with zipfile.ZipFile(output, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/media/padding.bin", b"\0" * 20_000)

    error = expect_error(ErrorCode.INVALID_FILE_TYPE, content=output.getvalue())

    assert "descomprimirse" in error.message


def test_check_file_accepts_a_real_workbook() -> None:
    check_file((FIXTURES / "validos.xlsx").read_bytes(), max_mb=2)


# ---------- Template (RF-08) ----------


def test_template_has_text_formatted_code_column_and_examples() -> None:
    book = load_workbook(io.BytesIO(build_template()))
    sheet = book.worksheets[0]

    assert TEMPLATE_FILENAME == "plantilla_codigos.xlsx"
    assert [sheet["A1"].value, sheet["B1"].value] == ["codigo", "nombre_archivo"]
    assert sheet["A2"].number_format == "@"
    assert sheet["A500"].number_format == "@"
    assert all(isinstance(sheet.cell(row, 1).value, str) for row in range(2, 5))
    assert book.sheetnames == ["Códigos", "Instrucciones"]


def test_template_examples_are_readable_and_valid() -> None:
    rows = read_sheet(build_template(), max_rows=10, max_mb=2, restore_leading_zeros=False)

    assert codes(rows) == ["786123456789", "7861234567898", "0012345678905"]
    assert all(row.error is None for row in rows)
